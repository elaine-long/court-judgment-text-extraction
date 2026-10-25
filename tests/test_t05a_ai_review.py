"""T05a: quote check, comparison rules, codebook-violation check, workbook write-back."""
import sys
from pathlib import Path

import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from src import ai_review as R  # noqa: E402


# ----------------------------------------------------------------------------- quotes
def test_quote_verbatim_and_na():
    src = "被告人李某，男，1969年6月10日出生于河南省睢县，汉族。"
    assert R.quote_ok("出生于河南省睢县", src)
    assert R.quote_ok("NA", src) and R.quote_ok("", src) and R.quote_ok(None, src)
    assert not R.quote_ok("被告人李某，男，出生于河南省睢县", src)      # stitched
    assert not R.quote_ok("可以从宽处理", src)                         # paraphrase


def test_quote_tolerates_encoding_damage_only():
    src = "1969年6月10日出生于河南���睢县，汉族"
    assert R.quote_ok("出生于河南省睢县", src)       # one character restored over a run of U+FFFD
    assert not R.quote_ok("出生于河北省睢县", src)    # a real change elsewhere still fails


def test_quote_length_is_separate():
    long_q = "判" * (R.QUOTE_MAX + 1)
    assert R.quote_too_long(long_q) and R.quote_ok(long_q, long_q)
    assert not R.quote_too_long("NA")


# ----------------------------------------------------------------------------- comparison
def test_agrees():
    assert R.agrees("term_months", "18", "18.0")
    assert R.agrees("term_months", "2.333", "2.333")
    assert not R.agrees("term_months", "NA", "3")
    assert R.agrees("procedure", "NA", float("nan"))           # empty prefill cell read by pandas
    assert R.agrees("origin_text", "贵州省荔波县", "荔波县")       # one contains the other
    assert not R.agrees("origin_text", "NA", "贵州省荔波县")
    assert not R.agrees("status_at_judgment", "在押", "NA")


def test_same_as_residence_uses_origin():
    assert R.SAME_AS.fullmatch("同上") and R.SAME_AS.fullmatch("住址同上") and R.SAME_AS.fullmatch("同户籍地")
    assert not R.SAME_AS.fullmatch("同安区")


def test_parse_validates_values():
    raw = ('{"origin_text": {"value": "NA", "quote": "NA"}, "residence_text": {"value": "NA", "quote": "NA"},'
           '"status_at_judgment": {"value": "在押", "quote": "现羁押于某看守所"}, "ever_bail": {"value": "0", "quote": "NA"},'
           '"plea_formal": {"value": "1", "quote": "认罪认罚"}, "jiejie": {"value": "0", "quote": "NA"},'
           '"procedure": {"value": "简易程序", "quote": "适用简易程序"}, "counsel_any": {"value": "0", "quote": "NA"},'
           '"penalty_type": {"value": "有期徒刑", "quote": "判处有期徒刑一年六个月"},'
           '"term_months": {"value": "18个月", "quote": "一年六个月"}, "probation": {"value": "0", "quote": "NA"}}')
    vals, quotes, errs = R.parse(raw)
    assert vals["term_months"] == "18"
    assert vals["procedure"] is None and "invalid:procedure" in errs   # not an allowed value


# ----------------------------------------------------------------------------- codebook violations
@pytest.mark.parametrize("value,quote,bad", [
    ("普通", "本院受理后，依法组成合议庭，公开开庭审理了本案", True),
    ("普通", "NA", True),
    ("简易", "本院依法适用简易程序，实行独任审判", False),
    ("普通", "转为普通程序", False),
    ("速裁", "（2019）浙0102刑初（速）123号", False),
    ("NA", "NA", False),
])
def test_procedure_violation(value, quote, bad):
    assert ("procedure" in R.rule_violations({"procedure": value}, {"procedure": quote})) is bad


@pytest.mark.parametrize("value,quote,bad", [
    ("在押", "即自2019年3月1日起至2019年9月1日止", True),
    ("在押", "判决执行以前先行羁押的，羁押一日折抵刑期一日", True),
    ("在押", "现羁押于某市看守所", False),
    ("在押", "2019年3月15日被逮捕", False),
    ("其他", "现正在某监狱服刑", False),
    ("NA", "NA", False),
])
def test_status_violation(value, quote, bad):
    assert ("status_at_judgment" in R.rule_violations({"status_at_judgment": value},
                                                       {"status_at_judgment": quote})) is bad


# ----------------------------------------------------------------------------- workbook write-back
def _toy_workbook(path):
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    for k, name in enumerate(("disputes", "audit")):
        ws = wb.active if k == 0 else wb.create_sheet(name)
        ws.title = name
        head = ["uid", "court_name", "judgment_date"]
        for f in R.WB_FIELDS:
            head += [f"{f}_prefill", f"{f}_final", f"{f}_evidence"]
        head += ["changed", "flag", "note", "minutes"]
        ws.append(head)
        for i in (1, 2):
            uid = f"{name[0]}{i:03d}"
            ws.append([uid, "某法院", "2019-01-01"] + [x for f in R.WB_FIELDS for x in ("NA", "NA", "ev")]
                      + [0, None, "keep me", 5])
        dv = DataValidation(type="list", formula1='"1,0,NA"')
        dv.add("H2:H3")  # nonlocal_origin_final
        ws.add_data_validation(dv)
        ws.freeze_panes = "D2"
    for name in ("codebook", "log"):
        wb.create_sheet(name).append(["x"])
    wb.save(path)


def _toy_cmp():
    rows = []
    for uid, blind in (("d001", 0), ("d002", 0), ("a001", 0), ("a002", 1)):
        r = {"uid": uid, "ai_blind": blind, "ai_disagree": "procedure" if uid == "d001" else "",
             "ai_quote_bad": "", "ai_rule_violation": "status_at_judgment" if uid == "a001" else ""}
        for f in R.WB_FIELDS:
            r[f"{f}_ai"], r[f"{f}_ai_quote"] = "1", "q"
        rows.append(r)
    return pd.DataFrame(rows)


def test_write_workbook_keeps_protected_cells(tmp_path):
    import T05a_ai_prereview as S
    from openpyxl import load_workbook
    src, dst, dst2 = tmp_path / "a.xlsx", tmp_path / "b.xlsx", tmp_path / "c.xlsx"
    _toy_workbook(src)
    S.write_workbook(src, dst, _toy_cmp())
    assert S.verify_protected(src, dst) > 0
    ws = load_workbook(dst)["disputes"]
    h = [c.value for c in ws[1]]
    assert h[h.index("origin_text_evidence") + 1: h.index("origin_text_evidence") + 3] == \
        ["origin_text_ai", "origin_text_ai_quote"]
    assert h[-4:] == ["ai_disagree", "ai_quote_bad", "ai_rule_violation", "ai_blind"]
    # the dropdown follows its column
    dv = ws.data_validations.dataValidation[0]
    assert str(dv.sqref) == f"{ws.cell(row=1, column=h.index('nonlocal_origin_final') + 1).column_letter}2:" \
                            f"{ws.cell(row=1, column=h.index('nonlocal_origin_final') + 1).column_letter}3"
    # blind row: AI columns empty, ai_blind = 1
    wa = load_workbook(dst)["audit"]
    ha = [c.value for c in wa[1]]
    row = [c.value for c in wa[3]]
    assert row[ha.index("ai_blind")] == 1
    assert all(row[j] is None for j, x in enumerate(ha) if x.endswith(("_ai", "_ai_quote")) or x in
               ("ai_disagree", "ai_rule_violation"))
    # re-running replaces the AI columns instead of adding more
    S.write_workbook(dst, dst2, _toy_cmp())
    assert [c.value for c in load_workbook(dst2)["disputes"][1]] == h
    assert S.verify_protected(src, dst2) > 0
